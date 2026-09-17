import os
from pathlib import Path

import pytest

from local_workspace_mcp.runner import PythonRunner

pytestmark = pytest.mark.skipif(not os.environ.get("LWMCP_DOCKER_TESTS"), reason="Opt-in real Docker test")


async def test_documents_and_isolation(tmp_path):
    (tmp_path / "sales.csv").write_text("month,revenue\nJan,100\nFeb,200\n")
    runner = PythonRunner(tmp_path, "local-workspace-mcp-worker:0.1.0")
    code = (Path(__file__).parent / "worker_smoke.py").read_text()
    result = await runner.run(code)
    assert result["exit_code"] == 0, result
    assert "VERIFIED total=300" in result["stdout"]
    assert {p.suffix for p in (tmp_path / "exports").iterdir()} >= {".docx", ".xlsx", ".pptx", ".pdf", ".png"}
    assert (tmp_path / "sales.csv").read_text() == "month,revenue\nJan,100\nFeb,200\n"


async def test_output_limit_and_container_cleanup(tmp_path):
    runner = PythonRunner(tmp_path, "local-workspace-mcp-worker:0.1.0")
    with pytest.raises(ValueError, match="output exceeded"):
        await runner.run("print('x' * 100000)")
    result = await runner.run("print('next job works')")
    assert result["exit_code"] == 0


async def test_office_rendering_and_recalculation(tmp_path):
    runner = PythonRunner(tmp_path, "local-workspace-mcp-worker:0.1.0")
    result = await runner.run("""
from docx import Document
from openpyxl import Workbook, load_workbook
from pathlib import Path
import subprocess
p=Document();p.add_paragraph('文件預覽測試');p.save('/output/preview.docx')
w=Workbook();w.active['A1']=100;w.active['A2']=200;w.active['A3']='=SUM(A1:A2)'
w.save('/tmp/calculation.xlsx')
for source, kind in [('/output/preview.docx','pdf'),('/tmp/calculation.xlsx','xlsx')]:
 r=subprocess.run(['libreoffice','-env:UserInstallation=file:///tmp/lo','--headless',
                   '--convert-to',kind,'--outdir','/output',source],capture_output=True,text=True)
 assert r.returncode == 0, r.stderr
assert load_workbook('/output/calculation.xlsx',data_only=True).active['A3'].value == 300
subprocess.run(['pdftoppm','-scale-to','800','-png','-singlefile',
                '/output/preview.pdf','/output/preview'],check=True)
assert Path('/output/preview.png').stat().st_size > 100
print('VERIFIED rendered Chinese DOCX and recalculated XLSX total=300')
""")
    assert result["exit_code"] == 0, result
    assert "recalculated XLSX total=300" in result["stdout"]
