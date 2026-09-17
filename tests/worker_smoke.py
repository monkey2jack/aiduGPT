"""Executed INSIDE the disposable worker by the integration test."""

import socket
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

sales = pd.read_csv("/workspace/sales.csv")
total = int(sales.revenue.sum())
assert total == 300
wb = Workbook()
ws = wb.active
ws.append(["Month", "Revenue"])
for row in sales.itertuples(index=False):
    ws.append(list(row))
ws.append(["Total", total])
wb.save("report.xlsx")
assert load_workbook("report.xlsx").active["B4"].value == 300

doc = Document()
doc.add_heading("Sales report", 0)
doc.add_paragraph(f"Total revenue: {total}")
doc.save("report.docx")
assert "300" in Document("report.docx").paragraphs[1].text

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Sales report"
slide.placeholders[1].text = f"Total revenue: {total}"
prs.save("report.pptx")
assert len(Presentation("report.pptx").slides) == 1

canvas = Canvas("report.pdf")
canvas.drawString(72, 720, f"Total revenue: {total}")
canvas.save()
assert "300" in PdfReader("report.pdf").pages[0].extract_text()

plt.bar(sales.month, sales.revenue)
plt.title("Revenue")
plt.savefig("chart.png")
assert Image.open("chart.png").width > 100

try:
    Path("/workspace/sales.csv").write_text("destroy")
    raise AssertionError("Input mount was writable")
except OSError:
    pass
try:
    socket.create_connection(("1.1.1.1", 443), timeout=2)
    raise AssertionError("Network was reachable")
except OSError:
    pass
assert not Path("/var/run/docker.sock").exists()
assert not Path("/workspace/../Users").exists()
print(f"VERIFIED total={total}; DOCX XLSX PPTX PDF PNG; read-only inputs; no network")
